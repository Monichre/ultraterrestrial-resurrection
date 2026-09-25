"""
Live Updates for Enhanced CocoIndex
Monitors file system changes and automatically updates the index
Based on https://github.com/cocoindex-io/cocoindex/tree/main/examples/live_updates
"""

import os
import asyncio
import logging
import hashlib
from typing import Dict, List, Optional, Callable, Any
from pathlib import Path
from datetime import datetime
import json
import time

# File monitoring
try:
    from watchdog.observers import Observer
    from watchdog.events import FileSystemEventHandler, FileModifiedEvent, FileCreatedEvent, FileDeletedEvent
    WATCHDOG_AVAILABLE = True
except ImportError:
    WATCHDOG_AVAILABLE = False
    print("Warning: watchdog not available. Install with: pip install watchdog")
    # Define dummy classes when watchdog is not available
    FileSystemEventHandler = object
    FileModifiedEvent = object
    FileCreatedEvent = object
    FileDeletedEvent = object
    Observer = None

from .backends.base import BackendInterface, Document


logger = logging.getLogger(__name__)


class FileChangeHandler(FileSystemEventHandler):
    """Handles file system events for live updates"""
    
    def __init__(self, live_updater: 'LiveUpdater'):
        self.live_updater = live_updater
        self.pending_changes = {}
        self.debounce_delay = 1.0  # Wait 1 second before processing changes
    
    def on_created(self, event):
        if not event.is_directory:
            self._handle_file_event(event.src_path, 'created')
    
    def on_modified(self, event):
        if not event.is_directory:
            self._handle_file_event(event.src_path, 'modified')
    
    def on_deleted(self, event):
        if not event.is_directory:
            self._handle_file_event(event.src_path, 'deleted')
    
    def _handle_file_event(self, file_path: str, event_type: str):
        """Handle file events with debouncing"""
        # Debounce rapid file changes
        current_time = time.time()
        if file_path in self.pending_changes:
            self.pending_changes[file_path]['last_modified'] = current_time
            self.pending_changes[file_path]['event_type'] = event_type
        else:
            self.pending_changes[file_path] = {
                'event_type': event_type,
                'last_modified': current_time
            }
        
        # Schedule processing after debounce delay
        asyncio.create_task(self._process_change_after_delay(file_path))
    
    async def _process_change_after_delay(self, file_path: str):
        """Process file change after debounce delay"""
        await asyncio.sleep(self.debounce_delay)
        
        if file_path in self.pending_changes:
            change_info = self.pending_changes.pop(file_path)
            await self.live_updater._process_file_change(file_path, change_info['event_type'])


class LiveUpdater:
    """Live updates manager for CocoIndex"""
    
    def __init__(self, 
                 backend: BackendInterface,
                 watch_directory: str,
                 file_extensions: Optional[List[str]] = None,
                 refresh_interval: float = 5.0,
                 content_extractor: Optional[Callable] = None,
                 metadata_extractor: Optional[Callable] = None):
        """
        Initialize live updater
        
        Args:
            backend: CocoIndex backend to update
            watch_directory: Directory to monitor for changes
            file_extensions: List of file extensions to monitor (e.g., ['.md', '.txt'])
            refresh_interval: Interval in seconds for periodic checks
            content_extractor: Function to extract content from files
            metadata_extractor: Function to extract metadata from files
        """
        self.backend = backend
        self.watch_directory = Path(watch_directory)
        self.file_extensions = file_extensions or ['.md', '.txt', '.pdf', '.docx']
        self.refresh_interval = refresh_interval
        self.content_extractor = content_extractor or self._default_content_extractor
        self.metadata_extractor = metadata_extractor or self._default_metadata_extractor
        
        # State tracking
        self.file_hashes = {}
        self.running = False
        self.observer = None
        self.stats = {
            'files_processed': 0,
            'files_added': 0,
            'files_updated': 0,
            'files_deleted': 0,
            'last_scan': None,
            'errors': 0
        }
        
        # Load existing state
        self._load_state()
    
    def _default_content_extractor(self, file_path: Path) -> str:
        """Default content extractor for text files"""
        try:
            if file_path.suffix.lower() == '.pdf':
                return self._extract_pdf_content(file_path)
            elif file_path.suffix.lower() in ['.docx']:
                return self._extract_docx_content(file_path)
            else:
                # Default text extraction
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
        except Exception as e:
            logger.error(f"Error extracting content from {file_path}: {e}")
            return ""
    
    def _extract_pdf_content(self, file_path: Path) -> str:
        """Extract content from PDF files"""
        try:
            import PyPDF2
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return text
        except ImportError:
            logger.warning("PyPDF2 not available for PDF extraction")
            return f"PDF file: {file_path.name}"
        except Exception as e:
            logger.error(f"Error extracting PDF content: {e}")
            return f"PDF file: {file_path.name}"
    
    def _extract_docx_content(self, file_path: Path) -> str:
        """Extract content from DOCX files"""
        try:
            import docx
            doc = docx.Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        except ImportError:
            logger.warning("python-docx not available for DOCX extraction")
            return f"DOCX file: {file_path.name}"
        except Exception as e:
            logger.error(f"Error extracting DOCX content: {e}")
            return f"DOCX file: {file_path.name}"
    
    def _default_metadata_extractor(self, file_path: Path) -> Dict[str, Any]:
        """Default metadata extractor"""
        try:
            stat = file_path.stat()
            return {
                'filename': file_path.name,
                'file_extension': file_path.suffix,
                'file_size': stat.st_size,
                'created_at': datetime.fromtimestamp(stat.st_ctime).isoformat(),
                'modified_at': datetime.fromtimestamp(stat.st_mtime).isoformat(),
                'file_path': str(file_path.absolute())
            }
        except Exception as e:
            logger.error(f"Error extracting metadata from {file_path}: {e}")
            return {'filename': file_path.name, 'error': str(e)}
    
    def _get_file_hash(self, file_path: Path) -> str:
        """Generate hash for file content"""
        try:
            with open(file_path, 'rb') as f:
                content = f.read()
                return hashlib.md5(content).hexdigest()
        except Exception as e:
            logger.error(f"Error hashing file {file_path}: {e}")
            return ""
    
    def _should_process_file(self, file_path: Path) -> bool:
        """Check if file should be processed based on extensions"""
        return file_path.suffix.lower() in self.file_extensions
    
    def _file_to_document_id(self, file_path: Path) -> str:
        """Generate document ID from file path"""
        # Use relative path from watch directory to create stable ID
        try:
            rel_path = file_path.relative_to(self.watch_directory)
            return str(rel_path).replace(os.sep, '/')
        except ValueError:
            # File is outside watch directory
            return str(file_path)
    
    async def _process_file_change(self, file_path: str, event_type: str):
        """Process a single file change"""
        try:
            file_path_obj = Path(file_path)
            
            if not self._should_process_file(file_path_obj):
                return
            
            doc_id = self._file_to_document_id(file_path_obj)
            
            if event_type == 'deleted':
                # Remove document from index
                success = await self.backend.delete_document(doc_id)
                if success:
                    self.stats['files_deleted'] += 1
                    if doc_id in self.file_hashes:
                        del self.file_hashes[doc_id]
                    logger.info(f"Deleted document: {doc_id}")
            
            elif event_type in ['created', 'modified']:
                # Check if file actually changed
                current_hash = self._get_file_hash(file_path_obj)
                if current_hash and current_hash != self.file_hashes.get(doc_id):
                    # Extract content and metadata
                    content = self.content_extractor(file_path_obj)
                    metadata = self.metadata_extractor(file_path_obj)
                    
                    # Create document
                    document = Document(
                        id=doc_id,
                        content=content,
                        metadata=metadata,
                        created_at=datetime.now().isoformat(),
                        updated_at=datetime.now().isoformat()
                    )
                    
                    # Check if document exists
                    existing_doc = await self.backend.get_document(doc_id)
                    
                    if existing_doc:
                        # Update existing document
                        success = await self.backend.update_document(doc_id, document)
                        if success:
                            self.stats['files_updated'] += 1
                            logger.info(f"Updated document: {doc_id}")
                    else:
                        # Add new document
                        success = await self.backend.add_documents([document])
                        if success:
                            self.stats['files_added'] += 1
                            logger.info(f"Added document: {doc_id}")
                    
                    if success:
                        self.file_hashes[doc_id] = current_hash
                        self.stats['files_processed'] += 1
            
            self._save_state()
            
        except Exception as e:
            logger.error(f"Error processing file change {file_path}: {e}")
            self.stats['errors'] += 1
    
    async def _scan_directory(self):
        """Perform full directory scan"""
        try:
            logger.info(f"Scanning directory: {self.watch_directory}")
            
            # Get all files in directory
            files_found = set()
            for file_path in self.watch_directory.rglob('*'):
                if file_path.is_file() and self._should_process_file(file_path):
                    files_found.add(self._file_to_document_id(file_path))
                    await self._process_file_change(str(file_path), 'modified')
            
            # Check for deleted files
            existing_docs = set(self.file_hashes.keys())
            deleted_files = existing_docs - files_found
            
            for doc_id in deleted_files:
                await self.backend.delete_document(doc_id)
                if doc_id in self.file_hashes:
                    del self.file_hashes[doc_id]
                self.stats['files_deleted'] += 1
                logger.info(f"Removed deleted file: {doc_id}")
            
            self.stats['last_scan'] = datetime.now().isoformat()
            self._save_state()
            
            logger.info(f"Directory scan completed. Processed {len(files_found)} files")
            
        except Exception as e:
            logger.error(f"Error during directory scan: {e}")
            self.stats['errors'] += 1
    
    def _save_state(self):
        """Save current state to disk"""
        try:
            state_file = self.watch_directory / '.cocoindex_live_state.json'
            state_data = {
                'file_hashes': self.file_hashes,
                'stats': self.stats,
                'last_saved': datetime.now().isoformat()
            }
            
            with open(state_file, 'w') as f:
                json.dump(state_data, f, indent=2)
                
        except Exception as e:
            logger.error(f"Error saving state: {e}")
    
    def _load_state(self):
        """Load state from disk"""
        try:
            state_file = self.watch_directory / '.cocoindex_live_state.json'
            if state_file.exists():
                with open(state_file, 'r') as f:
                    state_data = json.load(f)
                
                self.file_hashes = state_data.get('file_hashes', {})
                saved_stats = state_data.get('stats', {})
                
                # Merge with default stats
                for key, value in saved_stats.items():
                    if key in self.stats:
                        self.stats[key] = value
                
                logger.info(f"Loaded state: {len(self.file_hashes)} tracked files")
                
        except Exception as e:
            logger.error(f"Error loading state: {e}")
    
    async def start(self, initial_scan: bool = True):
        """Start live updates"""
        if self.running:
            logger.warning("Live updater is already running")
            return
        
        self.running = True
        logger.info(f"Starting live updates for {self.watch_directory}")
        
        # Perform initial scan
        if initial_scan:
            await self._scan_directory()
        
        # Start file watcher if available
        if WATCHDOG_AVAILABLE:
            try:
                handler = FileChangeHandler(self)
                self.observer = Observer()
                self.observer.schedule(handler, str(self.watch_directory), recursive=True)
                self.observer.start()
                logger.info("File watcher started")
            except Exception as e:
                logger.error(f"Error starting file watcher: {e}")
                self.observer = None
        
        # Start periodic scanning
        asyncio.create_task(self._periodic_scan())
    
    async def _periodic_scan(self):
        """Periodic directory scanning"""
        while self.running:
            try:
                await asyncio.sleep(self.refresh_interval)
                if self.running:
                    await self._scan_directory()
            except Exception as e:
                logger.error(f"Error in periodic scan: {e}")
                self.stats['errors'] += 1
    
    async def stop(self):
        """Stop live updates"""
        if not self.running:
            return
        
        self.running = False
        logger.info("Stopping live updates")
        
        # Stop file watcher
        if self.observer:
            self.observer.stop()
            self.observer.join()
            self.observer = None
        
        # Save final state
        self._save_state()
        
        logger.info("Live updates stopped")
    
    def get_stats(self) -> Dict[str, Any]:
        """Get live updater statistics"""
        return {
            **self.stats,
            'running': self.running,
            'watch_directory': str(self.watch_directory),
            'file_extensions': self.file_extensions,
            'refresh_interval': self.refresh_interval,
            'tracked_files': len(self.file_hashes),
            'watchdog_available': WATCHDOG_AVAILABLE,
            'observer_running': self.observer is not None and self.observer.is_alive() if self.observer else False
        }


class LiveCocoIndex:
    """Enhanced CocoIndex with live updates capability"""
    
    def __init__(self, 
                 backend: BackendInterface,
                 watch_directories: Optional[List[str]] = None,
                 **live_updater_kwargs):
        """
        Initialize LiveCocoIndex
        
        Args:
            backend: CocoIndex backend
            watch_directories: List of directories to monitor
            **live_updater_kwargs: Additional arguments for LiveUpdater
        """
        self.backend = backend
        self.watch_directories = watch_directories or []
        self.live_updaters: List[LiveUpdater] = []
        self.live_updater_kwargs = live_updater_kwargs
    
    async def initialize(self):
        """Initialize backend and live updaters"""
        # Initialize backend
        await self.backend.initialize()
        
        # Create live updaters for each directory
        for directory in self.watch_directories:
            updater = LiveUpdater(
                backend=self.backend,
                watch_directory=directory,
                **self.live_updater_kwargs
            )
            self.live_updaters.append(updater)
    
    async def start_live_updates(self, initial_scan: bool = True):
        """Start live updates for all directories"""
        for updater in self.live_updaters:
            await updater.start(initial_scan=initial_scan)
        
        logger.info(f"Started live updates for {len(self.live_updaters)} directories")
    
    async def stop_live_updates(self):
        """Stop live updates for all directories"""
        for updater in self.live_updaters:
            await updater.stop()
        
        logger.info("Stopped all live updates")
    
    def add_watch_directory(self, directory: str, start_immediately: bool = True):
        """Add a new directory to watch"""
        updater = LiveUpdater(
            backend=self.backend,
            watch_directory=directory,
            **self.live_updater_kwargs
        )
        self.live_updaters.append(updater)
        
        if start_immediately:
            asyncio.create_task(updater.start())
    
    def get_live_stats(self) -> Dict[str, Any]:
        """Get statistics from all live updaters"""
        stats = {
            'total_updaters': len(self.live_updaters),
            'updaters': []
        }
        
        for i, updater in enumerate(self.live_updaters):
            updater_stats = updater.get_stats()
            updater_stats['updater_id'] = i
            stats['updaters'].append(updater_stats)
        
        return stats
    
    # Delegate backend methods
    async def search(self, *args, **kwargs):
        return await self.backend.search(*args, **kwargs)
    
    async def add_documents(self, *args, **kwargs):
        return await self.backend.add_documents(*args, **kwargs)
    
    async def get_stats(self):
        return await self.backend.get_stats()
    
    # Add other backend methods as needed...


# Convenience function for easy setup
async def create_live_cocoindex(
    backend_type: str = "postgresql",
    watch_directories: Optional[List[str]] = None,
    **kwargs
) -> LiveCocoIndex:
    """
    Create a LiveCocoIndex with the specified backend
    
    Args:
        backend_type: Type of backend ('faiss' or 'postgresql')
        watch_directories: Directories to monitor
        **kwargs: Additional arguments for backend and live updater
    
    Returns:
        LiveCocoIndex: Configured live CocoIndex
    """
    from .backends.base import BackendFactory
    
    # Create backend
    backend = BackendFactory.create_backend(backend_type, **kwargs)
    
    # Create live CocoIndex
    live_cocoindex = LiveCocoIndex(
        backend=backend,
        watch_directories=watch_directories or [],
        **kwargs
    )
    
    # Initialize
    await live_cocoindex.initialize()
    
    return live_cocoindex