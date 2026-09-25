#!/usr/bin/env python3
"""
Bulk Folder Ingestion Script for Triple RAG System
Date: July 9, 2025 at 07:55 PST

Processes entire folders of PDF, text, and document files through the Triple RAG pipeline.
Designed specifically for ingesting large document collections like the Greer Document Library.
"""

import asyncio
import logging
import os
import sys
from pathlib import Path
from typing import Dict, List, Any, Optional, Set
from datetime import datetime
import json
import hashlib
import mimetypes
from dataclasses import dataclass
import traceback

# PDF and document processing
try:
    import PyPDF2
    import fitz  # PyMuPDF
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    import docx
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

# Add the lib directory to the path
sys.path.append(os.path.join(os.path.dirname(__file__), '..', 'lib'))

from integrations.triple_rag_integration import TripleRAGIntegration

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@dataclass
class DocumentFile:
    """Represents a document file for processing"""
    file_path: Path
    filename: str
    size_bytes: int
    mime_type: str
    file_hash: str
    last_modified: datetime
    supported: bool = True
    error: Optional[str] = None

class BulkFolderIngestion:
    """
    Handles bulk ingestion of document folders into Triple RAG system
    
    This class processes entire directories of PDF, text, and document files,
    extracting text content and routing everything through the Triple RAG pipeline.
    """
    
    # Supported file extensions
    SUPPORTED_EXTENSIONS = {
        '.pdf': 'PDF Document',
        '.txt': 'Text File',
        '.docx': 'Word Document',
        '.doc': 'Legacy Word Document',
        '.md': 'Markdown Document',
        '.rtf': 'Rich Text Format'
    }
    
    def __init__(self, folder_path: str, output_dir: str = "ingestion_results"):
        """
        Initialize bulk folder ingestion
        
        Args:
            folder_path: Path to folder containing documents
            output_dir: Directory for processing results and logs
        """
        self.folder_path = Path(folder_path)
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)
        
        # Processing statistics
        self.stats = {
            'total_files': 0,
            'supported_files': 0,
            'unsupported_files': 0,
            'processed_successfully': 0,
            'processing_failed': 0,
            'text_extraction_failed': 0,
            'total_pages_processed': 0,
            'total_text_length': 0,
            'start_time': None,
            'end_time': None,
            'processing_duration': None
        }
        
        # Error tracking
        self.errors = []
        self.warnings = []
        
        # Processed files tracking (avoid duplicates)
        self.processed_hashes: Set[str] = set()
        
        # Initialize Triple RAG integration
        self.triple_rag = TripleRAGIntegration()
        
        logger.info(f"Bulk folder ingestion initialized for: {folder_path}")
    
    def discover_documents(self) -> List[DocumentFile]:
        """
        Discover all document files in the target folder
        
        Returns:
            List of discovered document files
        """
        logger.info(f"Discovering documents in: {self.folder_path}")
        
        if not self.folder_path.exists():
            raise FileNotFoundError(f"Folder not found: {self.folder_path}")
        
        document_files = []
        
        # Walk through all files in directory (including subdirectories)
        for file_path in self.folder_path.rglob("*"):
            if file_path.is_file():
                self.stats['total_files'] += 1
                
                # Get file extension
                extension = file_path.suffix.lower()
                
                # Check if file type is supported
                if extension in self.SUPPORTED_EXTENSIONS:
                    try:
                        # Get file information
                        stat_info = file_path.stat()
                        file_size = stat_info.st_size
                        last_modified = datetime.fromtimestamp(stat_info.st_mtime)
                        
                        # Calculate file hash for duplicate detection
                        file_hash = self._calculate_file_hash(file_path)
                        
                        # Get MIME type
                        mime_type, _ = mimetypes.guess_type(str(file_path))
                        if not mime_type:
                            mime_type = f"application/{extension[1:]}"
                        
                        # Create document file object
                        doc_file = DocumentFile(
                            file_path=file_path,
                            filename=file_path.name,
                            size_bytes=file_size,
                            mime_type=mime_type,
                            file_hash=file_hash,
                            last_modified=last_modified,
                            supported=True
                        )
                        
                        document_files.append(doc_file)
                        self.stats['supported_files'] += 1
                        
                        logger.debug(f"Discovered: {file_path.name} ({file_size} bytes)")
                        
                    except Exception as e:
                        logger.warning(f"Error processing file info for {file_path.name}: {e}")
                        self.warnings.append(f"File info error for {file_path.name}: {e}")
                        
                        # Still add to list but mark as unsupported
                        doc_file = DocumentFile(
                            file_path=file_path,
                            filename=file_path.name,
                            size_bytes=0,
                            mime_type="unknown",
                            file_hash="",
                            last_modified=datetime.now(),
                            supported=False,
                            error=str(e)
                        )
                        document_files.append(doc_file)
                        self.stats['unsupported_files'] += 1
                else:
                    self.stats['unsupported_files'] += 1
                    logger.debug(f"Unsupported file type: {file_path.name} ({extension})")
        
        logger.info(f"Discovery complete: {len(document_files)} files found")
        logger.info(f"Supported: {self.stats['supported_files']}, Unsupported: {self.stats['unsupported_files']}")
        
        return document_files
    
    def _calculate_file_hash(self, file_path: Path) -> str:
        """
        Calculate SHA-256 hash of file for duplicate detection
        
        Args:
            file_path: Path to file
            
        Returns:
            SHA-256 hash string
        """
        sha256_hash = hashlib.sha256()
        try:
            with open(file_path, "rb") as f:
                # Read file in chunks to handle large files efficiently
                for chunk in iter(lambda: f.read(4096), b""):
                    sha256_hash.update(chunk)
            return sha256_hash.hexdigest()
        except Exception as e:
            logger.warning(f"Could not calculate hash for {file_path.name}: {e}")
            return f"error_{file_path.name}_{file_path.stat().st_size}"
    
    def extract_text_content(self, doc_file: DocumentFile) -> Dict[str, Any]:
        """
        Extract text content from document file
        
        Args:
            doc_file: Document file to process
            
        Returns:
            Extracted content information
        """
        logger.debug(f"Extracting text from: {doc_file.filename}")
        
        extraction_result = {
            'filename': doc_file.filename,
            'file_path': str(doc_file.file_path),
            'success': False,
            'text_content': '',
            'page_count': 0,
            'text_length': 0,
            'extraction_method': None,
            'error': None,
            'metadata': {
                'file_size': doc_file.size_bytes,
                'mime_type': doc_file.mime_type,
                'last_modified': doc_file.last_modified.isoformat(),
                'file_hash': doc_file.file_hash
            }
        }
        
        try:
            extension = doc_file.file_path.suffix.lower()
            
            if extension == '.pdf':
                extraction_result = self._extract_pdf_text(doc_file, extraction_result)
            elif extension == '.txt':
                extraction_result = self._extract_text_file(doc_file, extraction_result)
            elif extension == '.md':
                extraction_result = self._extract_text_file(doc_file, extraction_result)
            elif extension == '.docx':
                extraction_result = self._extract_docx_text(doc_file, extraction_result)
            elif extension == '.rtf':
                extraction_result = self._extract_text_file(doc_file, extraction_result)  # Basic RTF
            else:
                extraction_result['error'] = f"Unsupported file type: {extension}"
                return extraction_result
            
            # Update statistics
            if extraction_result['success']:
                self.stats['total_pages_processed'] += extraction_result['page_count']
                self.stats['total_text_length'] += extraction_result['text_length']
            else:
                self.stats['text_extraction_failed'] += 1
            
            return extraction_result
            
        except Exception as e:
            logger.error(f"Text extraction failed for {doc_file.filename}: {e}")
            extraction_result['error'] = str(e)
            self.stats['text_extraction_failed'] += 1
            return extraction_result
    
    def _extract_pdf_text(self, doc_file: DocumentFile, extraction_result: Dict[str, Any]) -> Dict[str, Any]:
        """Extract text from PDF file"""
        try:
            # Try PyMuPDF first (more reliable)
            if PYMUPDF_AVAILABLE:
                doc = fitz.open(str(doc_file.file_path))
                text_content = ""
                page_count = len(doc)
                
                for page_num in range(page_count):
                    page = doc.load_page(page_num)
                    text_content += page.get_text()
                    text_content += "\n\n"  # Page separator
                
                doc.close()
                
                extraction_result.update({
                    'success': True,
                    'text_content': text_content.strip(),
                    'page_count': page_count,
                    'text_length': len(text_content),
                    'extraction_method': 'PyMuPDF'
                })
                
            else:
                # Fallback to PyPDF2
                with open(doc_file.file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    text_content = ""
                    page_count = len(pdf_reader.pages)
                    
                    for page in pdf_reader.pages:
                        text_content += page.extract_text()
                        text_content += "\n\n"  # Page separator
                
                extraction_result.update({
                    'success': True,
                    'text_content': text_content.strip(),
                    'page_count': page_count,
                    'text_length': len(text_content),
                    'extraction_method': 'PyPDF2'
                })
                
        except Exception as e:
            extraction_result['error'] = f"PDF extraction failed: {e}"
            
        return extraction_result
    
    def _extract_text_file(self, doc_file: DocumentFile, extraction_result: Dict[str, Any]) -> Dict[str, Any]:
        """Extract text from plain text file"""
        try:
            # Try different encodings
            encodings = ['utf-8', 'utf-16', 'latin-1', 'cp1252']
            
            text_content = None
            encoding_used = None
            
            for encoding in encodings:
                try:
                    with open(doc_file.file_path, 'r', encoding=encoding) as file:
                        text_content = file.read()
                        encoding_used = encoding
                        break
                except UnicodeDecodeError:
                    continue
            
            if text_content is None:
                extraction_result['error'] = "Could not decode text file with any encoding"
                return extraction_result
            
            extraction_result.update({
                'success': True,
                'text_content': text_content,
                'page_count': 1,  # Text files are treated as single page
                'text_length': len(text_content),
                'extraction_method': f'text_file_{encoding_used}'
            })
            
        except Exception as e:
            extraction_result['error'] = f"Text file extraction failed: {e}"
            
        return extraction_result
    
    def _extract_docx_text(self, doc_file: DocumentFile, extraction_result: Dict[str, Any]) -> Dict[str, Any]:
        """Extract text from DOCX file"""
        try:
            if not PYTHON_DOCX_AVAILABLE:
                extraction_result['error'] = "python-docx library not available"
                return extraction_result
            
            doc = docx.Document(str(doc_file.file_path))
            text_content = ""
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                text_content += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text_content += cell.text + "\t"
                    text_content += "\n"
            
            extraction_result.update({
                'success': True,
                'text_content': text_content.strip(),
                'page_count': 1,  # DOCX doesn't have clear page boundaries
                'text_length': len(text_content),
                'extraction_method': 'python-docx'
            })
            
        except Exception as e:
            extraction_result['error'] = f"DOCX extraction failed: {e}"
            
        return extraction_result
    
    def create_triple_rag_document(self, doc_file: DocumentFile, extraction_result: Dict[str, Any]) -> Dict[str, Any]:
        """
        Create document structure for Triple RAG processing
        
        Args:
            doc_file: Source document file
            extraction_result: Text extraction results
            
        Returns:
            Document in Triple RAG format
        """
        # Generate unique document ID based on file hash
        doc_id = f"doc_{doc_file.file_hash[:16]}"
        
        # Create document title from filename
        title = doc_file.filename
        if doc_file.file_path.suffix:
            title = doc_file.file_path.stem  # Remove extension
        
        # Prepare metadata
        metadata = {
            'source': 'bulk_folder_ingestion',
            'source_type': 'local_file',
            'original_filename': doc_file.filename,
            'file_path': str(doc_file.file_path),
            'file_size': doc_file.size_bytes,
            'mime_type': doc_file.mime_type,
            'file_hash': doc_file.file_hash,
            'last_modified': doc_file.last_modified.isoformat(),
            'extraction_method': extraction_result.get('extraction_method'),
            'page_count': extraction_result.get('page_count', 0),
            'text_length': extraction_result.get('text_length', 0),
            'processed_at': datetime.now().isoformat(),
            'ingestion_batch': f"bulk_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        }
        
        # Create document structure
        document = {
            'id': doc_id,
            'title': title,
            'summary': extraction_result.get('text_content', ''),  # Full text as summary
            'url': f"file://{doc_file.file_path}",
            'metadata': json.dumps(metadata),
            'created_at': doc_file.last_modified.isoformat(),
            'source_type': 'local_file'
        }
        
        return document
    
    async def process_single_document(self, doc_file: DocumentFile) -> Dict[str, Any]:
        """
        Process a single document through the complete Triple RAG pipeline
        
        Args:
            doc_file: Document file to process
            
        Returns:
            Processing results
        """
        logger.info(f"Processing document: {doc_file.filename}")
        
        processing_result = {
            'filename': doc_file.filename,
            'file_path': str(doc_file.file_path),
            'doc_id': None,
            'status': 'pending',
            'text_extraction': {},
            'triple_rag_processing': {},
            'error': None,
            'processing_time': None,
            'timestamp': datetime.now().isoformat()
        }
        
        start_time = datetime.now()
        
        try:
            # Check for duplicate
            if doc_file.file_hash in self.processed_hashes:
                processing_result.update({
                    'status': 'skipped',
                    'error': 'Duplicate file (already processed)'
                })
                logger.info(f"Skipping duplicate: {doc_file.filename}")
                return processing_result
            
            # Extract text content
            logger.debug(f"Extracting text: {doc_file.filename}")
            extraction_result = self.extract_text_content(doc_file)
            processing_result['text_extraction'] = extraction_result
            
            if not extraction_result['success']:
                processing_result.update({
                    'status': 'failed',
                    'error': f"Text extraction failed: {extraction_result.get('error', 'unknown')}"
                })
                return processing_result
            
            # Create Triple RAG document
            logger.debug(f"Creating Triple RAG document: {doc_file.filename}")
            rag_document = self.create_triple_rag_document(doc_file, extraction_result)
            processing_result['doc_id'] = rag_document['id']
            
            # Process through Triple RAG pipeline
            logger.debug(f"Processing through Triple RAG: {doc_file.filename}")
            rag_result = await self.triple_rag.process_document(rag_document)
            processing_result['triple_rag_processing'] = rag_result
            
            if rag_result.get('status') == 'success':
                processing_result['status'] = 'completed'
                self.stats['processed_successfully'] += 1
                self.processed_hashes.add(doc_file.file_hash)
                logger.info(f"✅ Successfully processed: {doc_file.filename}")
            else:
                processing_result.update({
                    'status': 'failed',
                    'error': f"Triple RAG processing failed: {rag_result.get('error', 'unknown')}"
                })
                self.stats['processing_failed'] += 1
                logger.error(f"❌ Failed to process: {doc_file.filename}")
            
        except Exception as e:
            logger.error(f"Error processing {doc_file.filename}: {e}")
            logger.error(traceback.format_exc())
            processing_result.update({
                'status': 'error',
                'error': str(e)
            })
            self.stats['processing_failed'] += 1
            self.errors.append(f"Error processing {doc_file.filename}: {e}")
        
        # Calculate processing time
        end_time = datetime.now()
        processing_result['processing_time'] = (end_time - start_time).total_seconds()
        
        return processing_result
    
    async def process_all_documents(self, document_files: List[DocumentFile], batch_size: int = 10) -> List[Dict[str, Any]]:
        """
        Process all documents through Triple RAG pipeline
        
        Args:
            document_files: List of document files to process
            batch_size: Number of documents to process in parallel
            
        Returns:
            List of processing results
        """
        logger.info(f"Processing {len(document_files)} documents in batches of {batch_size}")
        
        # Initialize Triple RAG backends
        await self.triple_rag.initialize_backends()
        
        # Filter to only supported files
        supported_files = [doc for doc in document_files if doc.supported]
        logger.info(f"Processing {len(supported_files)} supported files")
        
        # Process in batches
        all_results = []
        
        for i in range(0, len(supported_files), batch_size):
            batch = supported_files[i:i + batch_size]
            batch_num = (i // batch_size) + 1
            total_batches = (len(supported_files) + batch_size - 1) // batch_size
            
            logger.info(f"Processing batch {batch_num}/{total_batches} ({len(batch)} files)")
            
            # Process batch in parallel
            batch_tasks = [self.process_single_document(doc_file) for doc_file in batch]
            batch_results = await asyncio.gather(*batch_tasks, return_exceptions=True)
            
            # Handle any exceptions
            for i, result in enumerate(batch_results):
                if isinstance(result, Exception):
                    logger.error(f"Exception in batch processing: {result}")
                    batch_results[i] = {
                        'filename': batch[i].filename,
                        'status': 'error',
                        'error': str(result),
                        'timestamp': datetime.now().isoformat()
                    }
            
            all_results.extend(batch_results)
            
            # Log batch completion
            successful_in_batch = sum(1 for r in batch_results if r.get('status') == 'completed')
            logger.info(f"Batch {batch_num} completed: {successful_in_batch}/{len(batch)} successful")
        
        return all_results
    
    async def run_bulk_ingestion(self, batch_size: int = 10, save_results: bool = True) -> Dict[str, Any]:
        """
        Run complete bulk folder ingestion process
        
        Args:
            batch_size: Number of documents to process in parallel
            save_results: Whether to save results to file
            
        Returns:
            Complete ingestion results
        """
        logger.info("🚀 Starting bulk folder ingestion")
        self.stats['start_time'] = datetime.now()
        
        try:
            # Step 1: Discover documents
            logger.info("📂 Discovering documents...")
            document_files = self.discover_documents()
            
            if not document_files:
                logger.warning("No documents found to process")
                return {
                    'status': 'completed',
                    'message': 'No documents found',
                    'statistics': self.stats,
                    'errors': self.errors,
                    'warnings': self.warnings
                }
            
            # Step 2: Process all documents
            logger.info("⚙️ Processing documents through Triple RAG...")
            processing_results = await self.process_all_documents(document_files, batch_size)
            
            # Step 3: Generate final results
            self.stats['end_time'] = datetime.now()
            self.stats['processing_duration'] = (self.stats['end_time'] - self.stats['start_time']).total_seconds()
            
            final_results = {
                'status': 'completed',
                'folder_path': str(self.folder_path),
                'processing_summary': {
                    'total_files_discovered': self.stats['total_files'],
                    'supported_files': self.stats['supported_files'],
                    'processed_successfully': self.stats['processed_successfully'],
                    'processing_failed': self.stats['processing_failed'],
                    'success_rate': (self.stats['processed_successfully'] / max(self.stats['supported_files'], 1)) * 100
                },
                'statistics': self.stats,
                'processing_results': processing_results,
                'triple_rag_statistics': self.triple_rag.get_statistics(),
                'errors': self.errors,
                'warnings': self.warnings,
                'timestamp': datetime.now().isoformat()
            }
            
            # Step 4: Save results if requested
            if save_results:
                results_file = self.output_dir / f"bulk_ingestion_results_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
                with open(results_file, 'w') as f:
                    json.dump(final_results, f, indent=2, default=str)
                logger.info(f"💾 Results saved to: {results_file}")
            
            # Step 5: Log final summary
            self._log_final_summary(final_results)
            
            return final_results
            
        except Exception as e:
            logger.error(f"Bulk ingestion failed: {e}")
            logger.error(traceback.format_exc())
            
            return {
                'status': 'failed',
                'error': str(e),
                'statistics': self.stats,
                'errors': self.errors + [str(e)],
                'warnings': self.warnings,
                'timestamp': datetime.now().isoformat()
            }
    
    def _log_final_summary(self, results: Dict[str, Any]):
        """Log final processing summary"""
        summary = results['processing_summary']
        stats = results['statistics']
        
        logger.info("📊 BULK INGESTION COMPLETE")
        logger.info("=" * 50)
        logger.info(f"📁 Folder: {results['folder_path']}")
        logger.info(f"📄 Total files discovered: {summary['total_files_discovered']}")
        logger.info(f"✅ Successfully processed: {summary['processed_successfully']}")
        logger.info(f"❌ Failed to process: {summary['processing_failed']}")
        logger.info(f"📈 Success rate: {summary['success_rate']:.1f}%")
        logger.info(f"📑 Total pages processed: {stats['total_pages_processed']}")
        logger.info(f"📝 Total text extracted: {stats['total_text_length']:,} characters")
        logger.info(f"⏱️ Processing duration: {stats['processing_duration']:.1f} seconds")
        
        if self.errors:
            logger.warning(f"⚠️ {len(self.errors)} errors occurred during processing")
        
        if self.warnings:
            logger.info(f"⚠️ {len(self.warnings)} warnings generated")

# CLI interface
async def main():
    """
    Main CLI interface for bulk folder ingestion
    """
    import argparse
    
    parser = argparse.ArgumentParser(description='Bulk folder ingestion for Triple RAG system')
    parser.add_argument('folder', help='Path to folder containing documents')
    parser.add_argument('--batch-size', type=int, default=10, help='Batch size for parallel processing')
    parser.add_argument('--output-dir', default='ingestion_results', help='Output directory for results')
    parser.add_argument('--no-save', action='store_true', help='Do not save results to file')
    
    args = parser.parse_args()
    
    print("🚀 Bulk Folder Ingestion for Triple RAG System")
    print("=" * 60)
    print(f"📁 Source folder: {args.folder}")
    print(f"📊 Batch size: {args.batch_size}")
    print(f"💾 Output directory: {args.output_dir}")
    print()
    
    # Initialize and run ingestion
    ingestion = BulkFolderIngestion(
        folder_path=args.folder,
        output_dir=args.output_dir
    )
    
    results = await ingestion.run_bulk_ingestion(
        batch_size=args.batch_size,
        save_results=not args.no_save
    )
    
    # Print final summary
    if results['status'] == 'completed':
        print("\n✅ Bulk ingestion completed successfully!")
        summary = results['processing_summary']
        print(f"Processed {summary['processed_successfully']}/{summary['supported_files']} files successfully")
        print(f"Success rate: {summary['success_rate']:.1f}%")
    else:
        print(f"\n❌ Bulk ingestion failed: {results.get('error', 'unknown error')}")

if __name__ == "__main__":
    asyncio.run(main())